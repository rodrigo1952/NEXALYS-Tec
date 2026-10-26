"""Inventario · Etapa 4: resumen, reportes, gráficos, alertas y Excel."""

import io
from datetime import date, timedelta
from decimal import Decimal

from django.urls import reverse
from django.utils import timezone

from catalogo.models import Producto
from catalogo.tests import BaseCatalogo

from . import reportes, servicios
from .models import FichaInventario, Venta
from .servicios import LineaPedida


class ReportesTests(BaseCatalogo):
    def setUp(self):
        super().setUp()
        self.laptop = self.crear_producto(nombre="Laptop", precio=Decimal("1000.00"), stock=10)
        FichaInventario.objects.create(producto=self.laptop, costo=Decimal("800.00"))
        self.mouse = self.crear_producto(nombre="Mouse", precio=Decimal("50.00"), stock=20)  # sin costo
        self.client.force_login(self.tienda.usuario)

    def vender(self, *lineas, dias_atras=0, metodo="efectivo"):
        venta = servicios.registrar_venta(self.tienda, self.tienda.usuario, [LineaPedida(*l) for l in lineas],
                                          metodo_pago=metodo)
        if dias_atras:
            Venta.objects.filter(pk=venta.pk).update(fecha=timezone.now() - timedelta(days=dias_atras))
        return venta

    def rango(self, **get):
        return reportes.rango_desde_get(get)[0]

    def test_periodos(self):
        hoy = timezone.localdate()
        self.assertEqual((self.rango().desde, self.rango().hasta), (hoy.replace(day=1), hoy))
        self.assertEqual(self.rango(periodo="7dias").desde, hoy - timedelta(days=6))
        anterior = self.rango(periodo="mes_anterior")
        self.assertEqual(anterior.hasta, hoy.replace(day=1) - timedelta(days=1))
        self.assertEqual(anterior.desde.day, 1)
        a_mano = self.rango(desde="2026-01-05", hasta="2026-01-20")
        self.assertEqual((a_mano.desde, a_mano.hasta, a_mano.clave), (date(2026, 1, 5), date(2026, 1, 20), ""))
        rango, error = reportes.rango_desde_get({"desde": "2026-02-10", "hasta": "2026-02-01"})
        self.assertIsNotNone(error)

    def test_totales_y_ganancia(self):
        self.vender((self.laptop.pk, 2), (self.mouse.pk, 4))
        anulada = self.vender((self.laptop.pk, 1))
        servicios.anular_venta(anulada, self.tienda.usuario, "Prueba de anulación")
        datos = reportes.resumen_ventas(self.tienda, self.rango(periodo="hoy"))
        self.assertEqual(datos["total"], Decimal("2200.00"))
        self.assertEqual((datos["cantidad"], datos["unidades"], datos["anuladas"]), (1, 6, 1))
        self.assertEqual(datos["ganancia"], Decimal("400.00"))  # solo la laptop tiene costo: 2 × (1000 − 800)
        self.assertEqual(datos["lineas_sin_costo"], 1)
        self.assertEqual(datos["margen"], Decimal("20.0"))
        self.assertEqual(datos["ticket"], Decimal("2200.00"))

    def test_las_ventas_fuera_del_periodo_no_cuentan(self):
        self.vender((self.laptop.pk, 1), dias_atras=40)
        self.assertEqual(reportes.resumen_ventas(self.tienda, self.rango(periodo="hoy"))["cantidad"], 0)

    def test_serie_por_dia_y_por_mes(self):
        self.vender((self.laptop.pk, 1))
        serie = reportes.ventas_en_el_tiempo(self.tienda, self.rango(periodo="7dias"))
        self.assertEqual(len(serie["etiquetas"]), 7)
        self.assertEqual(serie["vendido"][-1], 1000.0)
        self.assertEqual(serie["ganancia"][-1], 200.0)
        anual = reportes.ventas_en_el_tiempo(self.tienda, self.rango(desde="2025-01-01", hasta="2025-12-31"))
        self.assertTrue(anual["por_mes"])
        self.assertEqual(len(anual["etiquetas"]), 12)

    def test_mas_vendidos_y_metodos(self):
        self.vender((self.mouse.pk, 5), metodo="yape_plin")
        self.vender((self.laptop.pk, 1))
        top = reportes.mas_vendidos(self.tienda, self.rango(periodo="hoy"))
        self.assertEqual([f["producto_nombre"] for f in top], ["Mouse", "Laptop"])
        metodos = {m["metodo"]: m["suma"] for m in reportes.por_metodo_de_pago(self.tienda, self.rango(periodo="hoy"))}
        self.assertEqual(metodos, {"Efectivo": Decimal("1000.00"), "Yape / Plin": Decimal("250.00")})

    def test_inventario_valorizado(self):
        _, totales = reportes.inventario_valorizado(self.tienda)
        self.assertEqual(totales["valor_costo"], Decimal("8000.00"))      # 10 × 800 (el mouse no tiene costo)
        self.assertEqual(totales["valor_venta"], Decimal("11000.00"))     # 10 × 1000 + 20 × 50
        self.assertEqual(totales["sin_costo"], 1)

    def test_sin_movimiento(self):
        hace_mucho = timezone.now() - timedelta(days=100)
        Producto.objects.filter(pk__in=[self.laptop.pk, self.mouse.pk]).update(creado=hace_mucho)
        nuevo = self.crear_producto(nombre="Recién llegado", stock=3)  # creado hoy: no cuenta
        self.vender((self.laptop.pk, 1), dias_atras=10)  # se vendió hace poco
        self.vender((self.mouse.pk, 1), dias_atras=70)   # parado hace 70 días
        parados = reportes.sin_movimiento(self.tienda)
        self.assertEqual([p.nombre for p in parados], ["Mouse"])
        self.assertEqual(parados[0].dias_sin_vender, 70)
        self.assertNotIn(nuevo, parados)

    def test_alertas(self):
        FichaInventario.objects.filter(producto=self.laptop).update(stock_minimo=10)  # 10 ≤ 10: bajo
        self.crear_producto(nombre="Agotado", stock=0)
        self.assertEqual(reportes.cantidad_de_alertas(self.tienda), 2)
        self.assertEqual([p.nombre for p in reportes.alertas_de_stock(self.tienda)], ["Agotado", "Laptop"])

    def test_pagina_resumen(self):
        self.vender((self.laptop.pk, 1))
        respuesta = self.client.get(reverse("inventario:resumen"))
        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, "S/ 1,000.00")
        self.assertContains(respuesta, 'id="datos-grafico"')
        self.assertContains(respuesta, "chart.umd.min.js")
        self.assertContains(respuesta, reverse("inventario:resumen_excel"))
        self.assertContains(respuesta, f'href="{reverse("inventario:resumen")}"')  # pestaña

    def test_rango_invalido_avisa(self):
        respuesta = self.client.get(reverse("inventario:resumen"), {"desde": "2026-02-10", "hasta": "2026-02-01"})
        self.assertContains(respuesta, "no puede ser posterior")

    def test_circulito_de_alertas_en_el_menu(self):
        self.crear_producto(nombre="Agotado", stock=0)
        respuesta = self.client.get(reverse("inicio"))
        self.assertContains(respuesta, "alertas de stock</span>")
        self.client.logout()
        self.assertNotContains(self.client.get(reverse("inicio")), "alertas de stock</span>")

    def test_privado_de_cada_tienda(self):
        self.vender((self.laptop.pk, 1))
        self.client.force_login(self.otra_tienda.usuario)
        respuesta = self.client.get(reverse("inventario:resumen"))
        self.assertEqual(respuesta.context["ventas"]["cantidad"], 0)
        self.assertNotContains(respuesta, "Laptop")
        self.client.logout()
        self.assertNotEqual(self.client.get(reverse("inventario:resumen")).status_code, 200)

    def test_excel(self):
        from openpyxl import load_workbook

        self.vender((self.laptop.pk, 2), (self.mouse.pk, 1))
        respuesta = self.client.get(reverse("inventario:resumen_excel"), {"periodo": "hoy"})
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn("attachment;", respuesta["Content-Disposition"])
        self.assertIn(".xlsx", respuesta["Content-Disposition"])
        libro = load_workbook(io.BytesIO(respuesta.content))
        self.assertEqual(libro.sheetnames, ["Resumen", "Ventas", "Productos vendidos", "Inventario valorizado", "Sin movimiento"])
        resumen = {fila[0]: fila[1] for fila in libro["Resumen"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(resumen["Total vendido"], 2050)
        self.assertEqual(resumen["Ganancia"], 400)
        ventas = list(libro["Ventas"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(ventas), 1)
        self.assertEqual(ventas[0][-1], "Válida")

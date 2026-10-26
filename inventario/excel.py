"""
Descarga del resumen en Excel (Etapa 4): un archivo con varias hojas.

Usa openpyxl (está en requirements.txt). Las cantidades y montos se guardan como
números, no como texto, para que la tienda pueda sumarlos o graficarlos en Excel.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Venta
from .reportes import mas_vendidos

SOLES = '"S/" #,##0.00'
FECHA = "dd/mm/yyyy hh:mm"
ENCABEZADO = PatternFill("solid", fgColor="0B1D3A")  # azul marino de la marca


def _hoja(libro, titulo, columnas, filas, formatos=None, primera=False):
    hoja = libro.active if primera else libro.create_sheet()
    hoja.title = titulo
    hoja.append(columnas)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = ENCABEZADO
        celda.alignment = Alignment(vertical="center")
    for fila in filas:
        hoja.append(list(fila))
    for indice, formato in (formatos or {}).items():
        for (celda,) in hoja.iter_rows(min_row=2, min_col=indice, max_col=indice):
            celda.number_format = formato
    for indice, columna in enumerate(columnas, start=1):
        largo = max([len(str(columna))] + [len(str(c.value or "")) for (c,) in hoja.iter_rows(min_row=2, min_col=indice, max_col=indice)])
        hoja.column_dimensions[get_column_letter(indice)].width = min(max(largo + 2, 10), 60)
    hoja.freeze_panes = "A2"
    return hoja


def _sin_zona(fecha):
    """Excel no acepta fechas con zona horaria: se pasa a la hora local."""
    from django.utils import timezone
    return timezone.localtime(fecha).replace(tzinfo=None)


def libro_de_reportes(tienda, datos):
    rango, ventas, inventario = datos["rango"], datos["ventas"], datos["inventario"]
    libro = Workbook()

    resumen = [
        ("Tienda", tienda.nombre),
        ("Periodo", rango.etiqueta),
        ("Total vendido", ventas["total"]),
        ("Número de ventas", ventas["cantidad"]),
        ("Unidades vendidas", ventas["unidades"]),
        ("Ticket promedio", ventas["ticket"]),
        ("Ganancia", ventas["ganancia"]),
        ("Margen (%)", float(ventas["margen"]) if ventas["margen"] is not None else "—"),
        ("Líneas vendidas sin costo (no cuentan en la ganancia)", ventas["lineas_sin_costo"]),
        ("Ventas anuladas (no cuentan)", ventas["anuladas"]),
        ("Inventario: unidades", inventario["unidades"]),
        ("Inventario: valor a costo", inventario["valor_costo"]),
        ("Inventario: valor a precio de venta", inventario["valor_venta"]),
    ]
    hoja = _hoja(libro, "Resumen", ["Dato", "Valor"], resumen, primera=True)
    for fila in (4, 7, 8, 13, 14):
        hoja.cell(row=fila, column=2).number_format = SOLES

    lista = Venta.objects.filter(tienda=tienda, fecha__gte=rango.inicio(), fecha__lt=rango.fin()).order_by("fecha", "id")
    _hoja(libro, "Ventas", ["N.°", "Fecha", "Cliente", "DNI/RUC", "Método de pago", "Total", "Estado"], [
        (v.numero, _sin_zona(v.fecha), v.cliente_nombre or "Cliente varios", v.cliente_documento,
         v.get_metodo_pago_display(), v.total, "Anulada" if v.anulada else "Válida")
        for v in lista
    ], {2: FECHA, 6: SOLES})

    _hoja(libro, "Productos vendidos", ["Producto", "Unidades", "Vendido", "Ganancia"], [
        (f["producto_nombre"], f["unidades"], f["vendido"], f["ganancia"] if f["ganancia"] is not None else "—")
        for f in mas_vendidos(tienda, rango, limite=None)
    ], {3: SOLES, 4: SOLES})

    _hoja(libro, "Inventario valorizado",
          ["Producto", "Categoría", "Stock", "Precio de compra", "Valor a costo", "Precio de venta", "Valor a precio de venta"], [
              (p.nombre, p.categoria.nombre, p.stock, p.costo if p.costo is not None else "Falta",
               p.valor_costo if p.valor_costo is not None else "—", p.precio, p.valor_venta)
              for p in datos["productos_valorizados"]
          ], {4: SOLES, 5: SOLES, 6: SOLES, 7: SOLES})

    _hoja(libro, "Sin movimiento",
          ["Producto", "Stock", "Días sin vender", "Última venta", "Valor a costo"], [
              (p.nombre, p.stock, p.dias_sin_vender,
               _sin_zona(p.ultima_venta) if p.ultima_venta else "Nunca",
               p.valor_costo if p.valor_costo is not None else "—")
              for p in datos["sin_movimiento"]
          ], {4: FECHA, 5: SOLES})
    return libro
